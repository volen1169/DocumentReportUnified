
"""Centralized Export workspace for Hardware Assets and NAS Permissions."""

from __future__ import annotations

import html
import importlib
import io
import sys
from typing import Callable, Optional

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_ASSET_SOURCES = (
    ("Computer Asset", "Computer Asset", "💻"),
    ("Monitor Asset", "Asset Monitor", "🖥️"),
    ("Printer Asset", "Asset Printer", "🖨️"),
    ("Projector Asset", "Asset Projector", "📽️"),
    ("UPS Asset", "Asset UPS", "🔋"),
    ("CCTV Asset", "Asset CCTV", "📹"),
    ("Access Control Asset", "Asset Access Control", "🔐"),
)

_ALL_EXPORT_LABELS = tuple(item[0] for item in _ASSET_SOURCES) + ("NAS Permission",)


def _resolve_main_callable(name: str):
    """Resolve an existing application helper from Streamlit's main module."""
    main_module = sys.modules.get("__main__")
    if main_module is None:
        return None
    value = getattr(main_module, name, None)
    return value if callable(value) else None


def _load_nas_export_module():
    """Load the existing NAS export helper without coupling this view to one repo layout."""
    for module_name in ("services.nas_export", "nas_export"):
        try:
            return importlib.import_module(module_name)
        except Exception:
            continue
    return None


def _safe_frame(value) -> pd.DataFrame:
    return value.copy() if isinstance(value, pd.DataFrame) else pd.DataFrame()


def _clean_asset_export_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Remove internal transport-only columns while preserving the source schema."""
    frame = _safe_frame(frame)
    internal_columns = [
        column
        for column in frame.columns
        if str(column).startswith("@")
        or str(column) in {"_item_id", "id", "contentType", "fields@odata.context"}
    ]
    return frame.drop(columns=internal_columns, errors="ignore")


def _build_asset_csv(frame: pd.DataFrame) -> bytes:
    return _clean_asset_export_frame(frame).to_csv(index=False).encode("utf-8-sig")


def _style_asset_sheet(ws) -> None:
    """Apply a lightweight enterprise table format without changing any values."""
    if ws.max_column < 1:
        return

    thin = Side(style="thin", color="D7DEE8")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="1F2937")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx in range(1, ws.max_column + 1):
        values = [str(ws.cell(row=row_idx, column=col_idx).value or "") for row_idx in range(1, min(ws.max_row, 120) + 1)]
        width = min(max(max((len(value) for value in values), default=8) + 2, 11), 36)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24


def _build_asset_excel(frame: pd.DataFrame, sheet_name: str) -> bytes:
    frame = _clean_asset_export_frame(frame)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name=sheet_name[:31])
        _style_asset_sheet(writer.sheets[sheet_name[:31]])
    buffer.seek(0)
    return buffer.getvalue()


def _build_all_assets_excel(frames: dict[str, pd.DataFrame]) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for label, frame in frames.items():
            clean = _clean_asset_export_frame(frame)
            sheet_name = label.replace(" Asset", "")[:31]
            clean.to_excel(writer, index=False, sheet_name=sheet_name)
            _style_asset_sheet(writer.sheets[sheet_name])
    buffer.seek(0)
    return buffer.getvalue()


def _render_styles() -> None:
    st.html(
        """
        <style>
        .ex-header{
            display:flex;align-items:center;gap:15px;padding:18px 20px;margin:0 0 12px;
            background:#fff;border:1px solid #e2e8f0;border-radius:20px;
            box-shadow:0 10px 26px rgba(15,23,42,.05)
        }
        .ex-icon{
            width:54px;height:54px;display:grid;place-items:center;flex:0 0 54px;
            border-radius:16px;background:#eef2ff;color:#4f46e5;font-size:26px
        }
        .ex-header h1{margin:0;color:#0f172a;font-size:24px;font-weight:850}
        .ex-header p{margin:4px 0 0;color:#64748b;font-size:12px}
        .ex-grid{
            display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px;margin:10px 0 14px
        }
        .ex-source{
            min-height:88px;padding:13px 14px;background:#fff;border:1px solid #e2e8f0;
            border-radius:16px;box-shadow:0 6px 16px rgba(15,23,42,.035)
        }
        .ex-source b{display:block;color:#0f172a;font-size:12px}
        .ex-source span{display:block;margin-top:5px;color:#64748b;font-size:10px}
        .ex-source i{font-style:normal;font-size:20px}
        .ex-panel{
            padding:16px 17px;background:#fff;border:1px solid #e2e8f0;border-radius:18px;
            box-shadow:0 8px 20px rgba(15,23,42,.04)
        }
        .ex-panel-title{font-size:13px;font-weight:850;color:#0f172a;margin-bottom:3px}
        .ex-panel-sub{font-size:10.5px;color:#64748b;margin-bottom:12px}
        .ex-summary{
            display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:9px;margin:12px 0
        }
        .ex-stat{
            padding:11px 12px;border:1px solid #e5e7eb;border-radius:13px;background:#f8fafc
        }
        .ex-stat span{display:block;color:#64748b;font-size:9.5px}
        .ex-stat b{display:block;margin-top:3px;color:#0f172a;font-size:18px}
        @media(max-width:1100px){.ex-grid{grid-template-columns:repeat(2,minmax(0,1fr))}}
        @media(max-width:700px){.ex-grid,.ex-summary{grid-template-columns:1fr}}
        </style>
        """
    )


def render_report_view(
    *,
    load_sp_data: Optional[Callable] = None,
    load_nas_data: Optional[Callable] = None,
    clean_nas_principal: Optional[Callable] = None,
    nas_profile_lookup: Optional[Callable] = None,
) -> None:
    """
    Render the centralized Export page.

    Existing calls to render_report_view() remain valid. When callbacks are not
    supplied explicitly, the view reuses helpers already defined by the main app.
    """
    load_sp_data = load_sp_data or _resolve_main_callable("load_sp_data")
    load_nas_data = load_nas_data or _resolve_main_callable("load_nas_data")
    clean_nas_principal = clean_nas_principal or _resolve_main_callable("_clean_nas_principal")

    nas_module = _load_nas_export_module()

    if nas_profile_lookup is None:
        policy_lookup = _resolve_main_callable("get_user_internet_policy_summary")
        policy_formatter = _resolve_main_callable("format_policy_names")
        if (
            nas_module is not None
            and clean_nas_principal is not None
            and policy_lookup is not None
            and policy_formatter is not None
            and hasattr(nas_module, "resolve_nas_export_profile")
        ):
            nas_profile_lookup = lambda entity: nas_module.resolve_nas_export_profile(
                entity,
                clean_principal=clean_nas_principal,
                policy_lookup=policy_lookup,
                policy_formatter=policy_formatter,
            )

    _render_styles()

    st.markdown(
        """
        <section class="ex-header">
            <div class="ex-icon">⇩</div>
            <div>
                <h1>Export Center</h1>
                <p>ศูนย์กลางสำหรับ Export Hardware Asset และ NAS Permission</p>
            </div>
        </section>
        """,
        unsafe_allow_html=True,
    )

    cards = [
        ("💻", "Computer Asset", "SharePoint"),
        ("🖥️", "Monitor Asset", "SharePoint"),
        ("🖨️", "Printer Asset", "SharePoint"),
        ("📽️", "Projector Asset", "SharePoint"),
        ("🔋", "UPS Asset", "SharePoint"),
        ("📹", "CCTV Asset", "SharePoint"),
        ("🔐", "Access Control", "SharePoint"),
        ("🗂️", "NAS Permission", "NAS + AD"),
    ]
    st.markdown(
        '<div class="ex-grid">'
        + "".join(
            f'<div class="ex-source"><i>{icon}</i><b>{html.escape(name)}</b><span>{html.escape(source)}</span></div>'
            for icon, name, source in cards
        )
        + "</div>",
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="ex-panel"><div class="ex-panel-title">เลือกข้อมูลที่ต้องการ Export</div>'
        '<div class="ex-panel-sub">โหลดเฉพาะชุดข้อมูลที่เลือก เพื่อลดเวลาในการเปิดหน้า Export</div></div>',
        unsafe_allow_html=True,
    )

    selected = st.selectbox(
        "ประเภทข้อมูล",
        _ALL_EXPORT_LABELS,
        key="export_center_source",
    )

    if selected == "NAS Permission":
        if load_nas_data is None:
            st.error("ไม่พบ load_nas_data จากแอปหลัก")
            return
        if nas_module is None:
            st.error("ไม่พบโมดูล NAS export (รองรับ services.nas_export หรือ nas_export)")
            return
        if clean_nas_principal is None or nas_profile_lookup is None:
            st.error("ยังไม่พร้อม Export NAS: ไม่พบ clean principal / AD profile lookup")
            return

        with st.spinner("กำลังเตรียม NAS Permission..."):
            nas_source = _safe_frame(load_nas_data())

        if nas_source.empty:
            st.info("ไม่พบข้อมูล NAS Permission สำหรับ Export")
            return

        if not hasattr(nas_module, "build_nas_export_dataframe"):
            st.error("NAS export module ไม่มี build_nas_export_dataframe")
            return

        export_df = nas_module.build_nas_export_dataframe(
            nas_source,
            clean_principal=clean_nas_principal,
            profile_lookup=nas_profile_lookup,
        )

        st.markdown(
            f'<div class="ex-summary">'
            f'<div class="ex-stat"><span>Users</span><b>{len(export_df):,}</b></div>'
            f'<div class="ex-stat"><span>Columns</span><b>{len(export_df.columns):,}</b></div>'
            f'<div class="ex-stat"><span>Source Shares</span><b>{nas_source["Share"].nunique() if "Share" in nas_source else 0:,}</b></div>'
            f'</div>',
            unsafe_allow_html=True,
        )

        preview_rows = min(len(export_df), 20)
        st.dataframe(export_df.head(preview_rows), use_container_width=True, hide_index=True)

        c1, c2 = st.columns(2)
        with c1:
            st.download_button(
                "⬇️ Export NAS Permission.xlsx",
                data=nas_module.build_nas_excel(export_df),
                file_name="NAS_Permission.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="primary",
                key="export_nas_excel",
            )
        with c2:
            st.download_button(
                "⬇️ Export NAS Permission.csv",
                data=nas_module.build_nas_csv(export_df),
                file_name="NAS_Permission.csv",
                mime="text/csv",
                use_container_width=True,
                key="export_nas_csv",
            )
        return

    if load_sp_data is None:
        st.error("ไม่พบ load_sp_data จากแอปหลัก")
        return

    source_lookup = {label: list_name for label, list_name, _ in _ASSET_SOURCES}
    list_name = source_lookup[selected]

    with st.spinner(f"กำลังโหลด {selected}..."):
        frame = _safe_frame(load_sp_data(list_name))

    clean_frame = _clean_asset_export_frame(frame)

    st.markdown(
        f'<div class="ex-summary">'
        f'<div class="ex-stat"><span>Dataset</span><b>{html.escape(selected)}</b></div>'
        f'<div class="ex-stat"><span>Records</span><b>{len(clean_frame):,}</b></div>'
        f'<div class="ex-stat"><span>Columns</span><b>{len(clean_frame.columns):,}</b></div>'
        f'</div>',
        unsafe_allow_html=True,
    )

    if clean_frame.empty:
        st.info(f"ไม่พบข้อมูล {selected}")
    else:
        st.dataframe(clean_frame.head(20), use_container_width=True, hide_index=True)

    c1, c2 = st.columns(2)
    with c1:
        st.download_button(
            f"⬇️ Export {selected}.xlsx",
            data=_build_asset_excel(clean_frame, selected),
            file_name=f"{selected.replace(' ', '_')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            type="primary",
            key="export_asset_excel",
        )
    with c2:
        st.download_button(
            f"⬇️ Export {selected}.csv",
            data=_build_asset_csv(clean_frame),
            file_name=f"{selected.replace(' ', '_')}.csv",
            mime="text/csv",
            use_container_width=True,
            key="export_asset_csv",
        )

    st.divider()
    with st.expander("Export Hardware Asset ทั้งหมดเป็น Excel เดียว", expanded=False):
        st.caption("สร้าง Workbook 7 Sheet: Computer, Monitor, Printer, Projector, UPS, CCTV และ Access Control")
        if st.button("เตรียม Hardware Asset Workbook", use_container_width=True, key="prepare_all_asset_export"):
            frames = {}
            progress = st.progress(0)
            for index, (label, asset_list_name, _) in enumerate(_ASSET_SOURCES, start=1):
                frames[label] = _safe_frame(load_sp_data(asset_list_name))
                progress.progress(index / len(_ASSET_SOURCES))
            workbook = _build_all_assets_excel(frames)
            st.session_state["all_asset_export_workbook"] = workbook
            st.session_state["all_asset_export_counts"] = {
                label: len(frame) for label, frame in frames.items()
            }

        workbook = st.session_state.get("all_asset_export_workbook")
        if workbook:
            counts = st.session_state.get("all_asset_export_counts", {})
            st.success(
                "พร้อมดาวน์โหลด — "
                + " | ".join(f"{label}: {count:,}" for label, count in counts.items())
            )
            st.download_button(
                "⬇️ Export All Hardware Assets.xlsx",
                data=workbook,
                file_name="Hardware_Assets_All.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="primary",
                key="download_all_assets",
            )
