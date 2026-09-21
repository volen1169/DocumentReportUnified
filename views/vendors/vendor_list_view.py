"""Vendor Directory backed by VendorList.XLSX on SharePoint."""

from __future__ import annotations

import html
import io
from copy import copy
from datetime import date, datetime

import pandas as pd
import requests
import streamlit as st
from openpyxl import load_workbook

from services.excel_storage import SHAREPOINT_FOLDER, _get_sharepoint_drive_id
from services.microsoft_graph import GRAPH_URL, get_access_token, get_sp_site_id


VENDOR_FILE_NAME = "VendorList.XLSX"
VENDOR_WORKBOOK_PATH = f"{SHAREPOINT_FOLDER}/{VENDOR_FILE_NAME}"
INACTIVE_KEYWORDS = ("ยกเลิกการใช้แล้ว", "ยกเลิก", "inactive", "เลิกใช้")


def _clean(value) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _normalise_header(value) -> str:
    return " ".join(_clean(value).lower().replace("-", " ").replace("_", " ").split())


def _find_column(headers, *candidates):
    lookup = {_normalise_header(header): header for header in headers}
    for candidate in candidates:
        match = lookup.get(_normalise_header(candidate))
        if match:
            return match
    return None


def vendor_status(row, status_column=None, remark_column=None) -> str:
    if status_column:
        raw_status = _clean(row.get(status_column)).lower()
        if raw_status:
            return "Inactive" if any(keyword in raw_status for keyword in INACTIVE_KEYWORDS) else "Active"
    remark = _clean(row.get(remark_column)).lower() if remark_column else ""
    return "Inactive" if any(keyword in remark for keyword in INACTIVE_KEYWORDS) else "Active"


def _worksheet_headers(worksheet):
    headers = []
    for column in range(1, worksheet.max_column + 1):
        value = worksheet.cell(1, column).value
        if value is None or not str(value).strip():
            continue
        headers.append(str(value).strip())
    return headers


def _select_vendor_worksheet(workbook):
    if "Vendor" in workbook.sheetnames:
        return workbook["Vendor"]
    for worksheet in workbook.worksheets:
        if _worksheet_headers(worksheet):
            return worksheet
    raise ValueError("ไม่พบ worksheet ที่มี Header ใน VendorList.XLSX")


def _download_vendor_workbook():
    token = get_access_token()
    if not token:
        raise RuntimeError("ไม่สามารถรับ Microsoft Graph access token ได้")
    site_id = get_sp_site_id()
    if not site_id:
        raise RuntimeError("ไม่พบ SharePoint Site ID")
    drive_id = _get_sharepoint_drive_id()
    response = requests.get(
        f"{GRAPH_URL}/drives/{drive_id}/root:/{VENDOR_WORKBOOK_PATH}:/content",
        headers={"Authorization": f"Bearer {token}"},
        timeout=45,
    )
    response.raise_for_status()
    return load_workbook(io.BytesIO(response.content)), drive_id


@st.cache_data(ttl=180)
def load_vendor_workbook():
    """Return serialisable Vendor data; writable workbooks are always re-downloaded."""
    workbook, _ = _download_vendor_workbook()
    worksheet = _select_vendor_worksheet(workbook)
    headers = _worksheet_headers(worksheet)
    status_column = _find_column(headers, "Status")
    remark_column = _find_column(headers, "Remark")
    records = []
    for row_number in range(2, worksheet.max_row + 1):
        record = {
            header: worksheet.cell(row_number, column_number).value
            for column_number, header in enumerate(headers, start=1)
        }
        if not any(_clean(value) for value in record.values()):
            continue
        record["__worksheet"] = worksheet.title
        record["__source_row"] = row_number
        record["__status"] = vendor_status(record, status_column, remark_column)
        records.append(record)
    return {
        "worksheet": worksheet.title,
        "headers": headers,
        "records": records,
        "site_id": get_sp_site_id(),
        "path": VENDOR_WORKBOOK_PATH,
    }


def upload_vendor_workbook(workbook):
    token = get_access_token()
    if not token:
        return False, "ไม่สามารถรับ Microsoft Graph access token ได้"
    drive_id = _get_sharepoint_drive_id()
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    response = requests.put(
        f"{GRAPH_URL}/drives/{drive_id}/root:/{VENDOR_WORKBOOK_PATH}:/content",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/octet-stream"},
        data=output.getvalue(),
        timeout=45,
    )
    if response.status_code not in (200, 201):
        return False, f"HTTP {response.status_code}: {response.text[:300]}"
    load_vendor_workbook.clear()
    return True, "บันทึกสำเร็จ"


def filter_vendor_records(records, headers, search="", vendor_type="All", status="All"):
    company_column = _find_column(headers, "Company/Shop", "Company", "Shop")
    type_column = _find_column(headers, "Type", "Category")
    searchable = [
        column
        for column in (
            company_column,
            type_column,
            _find_column(headers, "Contact name", "Contact", "Contact Person"),
            _find_column(headers, "Telephone number", "Telephone", "Phone"),
            _find_column(headers, "E-mail", "Email"),
            _find_column(headers, "Remark", "Remarks"),
        )
        if column
    ]
    query = _clean(search).casefold()
    filtered = []
    for record in records:
        if query and not any(query in _clean(record.get(column)).casefold() for column in searchable):
            continue
        if vendor_type != "All" and (not type_column or _clean(record.get(type_column)) != vendor_type):
            continue
        if status != "All" and record.get("__status") != status:
            continue
        filtered.append(record)
    return sorted(
        filtered,
        key=lambda record: (
            0 if status == "Inactive" or record.get("__status") == "Active" else 1,
            _clean(record.get(company_column)).casefold() if company_column else "",
        ),
    )


def _copy_row_style(worksheet, source_row, target_row):
    if source_row < 2:
        return
    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    if source_row in worksheet.row_dimensions:
        worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height


def _resize_tables(worksheet):
    if not worksheet.tables:
        return
    last_column = worksheet.max_column
    while last_column > 1 and not _clean(worksheet.cell(1, last_column).value):
        last_column -= 1
    from openpyxl.utils import get_column_letter

    table_ref = f"A1:{get_column_letter(last_column)}{max(1, worksheet.max_row)}"
    for table in worksheet.tables.values():
        table.ref = table_ref


def _input_for_header(header, value=None, key=None):
    label = str(header)
    normalised = _normalise_header(header)
    if normalised == "date":
        parsed = value if isinstance(value, (date, datetime)) else None
        if isinstance(parsed, datetime):
            parsed = parsed.date()
        return st.date_input(label, value=parsed or date.today(), key=key)
    if normalised in ("remark", "remarks", "address"):
        return st.text_area(label, value=_clean(value), key=key)
    return st.text_input(label, value=_clean(value), key=key)


@st.dialog("Vendor Details")
def _view_vendor_dialog(record, headers):
    for header in headers:
        value = _clean(record.get(header)) or "-"
        st.markdown(f"**{html.escape(str(header))}**")
        st.write(value)


@st.dialog("+ Add Vendor")
def _add_vendor_dialog(headers):
    values = {header: _input_for_header(header, key=f"vendor_add_{index}") for index, header in enumerate(headers)}
    company_column = _find_column(headers, "Company/Shop", "Company", "Shop")
    if st.button("Save Vendor", type="primary", use_container_width=True):
        if company_column and not _clean(values.get(company_column)):
            st.error(f"{company_column} ห้ามว่าง")
            return
        try:
            workbook, _ = _download_vendor_workbook()
            worksheet = _select_vendor_worksheet(workbook)
            latest_headers = _worksheet_headers(worksheet)
            next_row = worksheet.max_row + 1
            _copy_row_style(worksheet, worksheet.max_row, next_row)
            for column_number, header in enumerate(latest_headers, start=1):
                worksheet.cell(next_row, column_number, values.get(header) or None)
            _resize_tables(worksheet)
            ok, message = upload_vendor_workbook(workbook)
            if not ok:
                st.error(f"เพิ่ม Vendor ไม่สำเร็จ: {message}")
                return
            st.success("เพิ่ม Vendor สำเร็จ")
            st.rerun()
        except Exception as error:
            st.error(f"เพิ่ม Vendor ไม่สำเร็จ: {error}")


@st.dialog("Edit Vendor")
def _edit_vendor_dialog(record, headers):
    source_sheet = record.get("__worksheet")
    source_row = int(record.get("__source_row"))
    values = {
        header: _input_for_header(header, record.get(header), key=f"vendor_edit_{source_sheet}_{source_row}_{index}")
        for index, header in enumerate(headers)
    }
    company_column = _find_column(headers, "Company/Shop", "Company", "Shop")
    save_column, delete_column = st.columns(2)
    with save_column:
        save_clicked = st.button("Save Changes", type="primary", use_container_width=True)
    with delete_column:
        delete_clicked = st.button("Delete", use_container_width=True)
    confirm_key = f"vendor_delete_confirm_{source_sheet}_{source_row}"
    if delete_clicked:
        st.session_state[confirm_key] = True
    if st.session_state.get(confirm_key):
        st.warning("ยืนยันการลบ Vendor รายการนี้หรือไม่?")
        if st.button("Confirm Delete", type="primary", use_container_width=True):
            try:
                workbook, _ = _download_vendor_workbook()
                if source_sheet not in workbook.sheetnames:
                    raise ValueError(f"ไม่พบ worksheet: {source_sheet}")
                worksheet = workbook[source_sheet]
                if source_row < 2 or source_row > worksheet.max_row:
                    raise ValueError("ไม่พบ source row ที่ต้องการลบ")
                worksheet.delete_rows(source_row, 1)
                _resize_tables(worksheet)
                ok, message = upload_vendor_workbook(workbook)
                if not ok:
                    st.error(f"ลบ Vendor ไม่สำเร็จ: {message}")
                    return
                st.session_state.pop(confirm_key, None)
                st.success("ลบ Vendor สำเร็จ")
                st.rerun()
            except Exception as error:
                st.error(f"ลบ Vendor ไม่สำเร็จ: {error}")
    if save_clicked:
        if company_column and not _clean(values.get(company_column)):
            st.error(f"{company_column} ห้ามว่าง")
            return
        try:
            workbook, _ = _download_vendor_workbook()
            if source_sheet not in workbook.sheetnames:
                raise ValueError(f"ไม่พบ worksheet: {source_sheet}")
            worksheet = workbook[source_sheet]
            latest_headers = _worksheet_headers(worksheet)
            if source_row < 2 or source_row > worksheet.max_row:
                raise ValueError("ไม่พบ source row ที่ต้องการแก้ไข")
            for column_number, header in enumerate(latest_headers, start=1):
                worksheet.cell(source_row, column_number, values.get(header) or None)
            ok, message = upload_vendor_workbook(workbook)
            if not ok:
                st.error(f"แก้ไข Vendor ไม่สำเร็จ: {message}")
                return
            st.success("แก้ไข Vendor สำเร็จ")
            st.rerun()
        except Exception as error:
            st.error(f"แก้ไข Vendor ไม่สำเร็จ: {error}")


def _metric_card(label, value, accent="#4F46E5"):
    st.markdown(
        f"""<div style="background:#fff;border:1px solid #E2E8F0;border-radius:16px;padding:16px;box-shadow:0 4px 14px rgba(15,23,42,.05);border-top:3px solid {accent}">
        <div style="color:#64748B;font-size:.82rem">{html.escape(label)}</div><div style="font-size:1.75rem;font-weight:750;color:#0F172A">{value}</div></div>""",
        unsafe_allow_html=True,
    )


def _vendor_card(record, headers, admin_mode, key_prefix):
    company_column = _find_column(headers, "Company/Shop", "Company", "Shop")
    type_column = _find_column(headers, "Type", "Category")
    contact_column = _find_column(headers, "Contact name", "Contact", "Contact Person")
    telephone_column = _find_column(headers, "Telephone number", "Telephone", "Phone")
    status = record.get("__status", "Active")
    inactive = status == "Inactive"
    opacity = ".68" if inactive else "1"
    badge_bg = "#FEE2E2" if inactive else "#DCFCE7"
    badge_color = "#991B1B" if inactive else "#166534"
    title = html.escape(_clean(record.get(company_column)) or "Unnamed Vendor")
    vendor_type = html.escape(_clean(record.get(type_column)) or "-")
    contact = html.escape(_clean(record.get(contact_column)) or "-")
    telephone = html.escape(_clean(record.get(telephone_column)) or "-")
    st.markdown(
        f"""<div style="min-height:190px;background:#fff;border:1px solid #E2E8F0;border-radius:18px;padding:18px;box-shadow:0 5px 16px rgba(15,23,42,.05);opacity:{opacity}">
        <div style="font-weight:750;font-size:1.02rem;color:#0F172A">🏢 {title}</div>
        <div style="display:inline-block;margin:8px 0 14px;padding:3px 9px;border-radius:999px;background:#EEF2FF;color:#4338CA;font-size:.78rem">{vendor_type}</div>
        <div style="color:#475569;margin-bottom:5px">👤 {contact}</div><div style="color:#475569">📞 {telephone}</div>
        <div style="display:inline-block;margin-top:14px;padding:3px 9px;border-radius:999px;background:{badge_bg};color:{badge_color};font-size:.76rem;font-weight:700">{status}</div></div>""",
        unsafe_allow_html=True,
    )
    view_column, edit_column = st.columns(2)
    with view_column:
        if st.button("View Details", key=f"vendor_view_{key_prefix}", use_container_width=True):
            _view_vendor_dialog(record, headers)
    with edit_column:
        if admin_mode and st.button("Edit", key=f"vendor_edit_{key_prefix}", use_container_width=True):
            _edit_vendor_dialog(record, headers)


def render_vendor_list_page(admin_mode=False):
    st.markdown("## Vendor Directory")
    st.caption("Supplier & Contact Management")
    try:
        with st.spinner("กำลังโหลด VendorList.XLSX จาก SharePoint..."):
            payload = load_vendor_workbook()
    except Exception as error:
        st.error("ไม่สามารถโหลด VendorList.XLSX จาก SharePoint ได้")
        st.caption(str(error))
        return
    headers = payload.get("headers", [])
    records = payload.get("records", [])
    type_column = _find_column(headers, "Type", "Category")
    active_count = sum(record.get("__status") == "Active" for record in records)
    inactive_count = len(records) - active_count
    type_values = sorted({_clean(record.get(type_column)) for record in records if type_column and _clean(record.get(type_column))}, key=str.casefold)
    metrics = st.columns(4)
    with metrics[0]:
        _metric_card("Total Vendors", len(records))
    with metrics[1]:
        _metric_card("Active Vendors", active_count, "#16A34A")
    with metrics[2]:
        _metric_card("Inactive Vendors", inactive_count, "#94A3B8")
    with metrics[3]:
        _metric_card("Categories / Types", len(type_values), "#2563EB")
    st.write("")
    search_column, type_filter_column, status_filter_column, add_column = st.columns([2.2, 1.2, 1.1, 1.1])
    with search_column:
        search = st.text_input("Search", placeholder="🔎 Search vendor...", label_visibility="collapsed", key="vendor_search")
    with type_filter_column:
        selected_type = st.selectbox("Type", ["All", *type_values], key="vendor_type_filter")
    with status_filter_column:
        selected_status = st.selectbox("Status", ["All", "Active", "Inactive"], key="vendor_status_filter")
    with add_column:
        if admin_mode and st.button("+ Add Vendor", type="primary", use_container_width=True):
            _add_vendor_dialog(headers)
    filter_signature = (search.strip().casefold(), selected_type, selected_status)
    if st.session_state.get("vendor_filter_signature") != filter_signature:
        st.session_state.vendor_filter_signature = filter_signature
        st.session_state.vendor_page = 1
    filtered = filter_vendor_records(records, headers, search, selected_type, selected_status)
    if not filtered:
        st.info("ยังไม่มีข้อมูล Vendor")
        if admin_mode and st.button("+ Add Vendor", key="vendor_empty_add", type="primary"):
            _add_vendor_dialog(headers)
        return
    page_size = st.selectbox("รายการต่อหน้า", [12, 24, 48], key="vendor_page_size")
    total_pages = max(1, (len(filtered) + page_size - 1) // page_size)
    current_page = min(max(1, int(st.session_state.get("vendor_page", 1))), total_pages)
    st.session_state.vendor_page = current_page
    start = (current_page - 1) * page_size
    page_records = filtered[start : start + page_size]
    st.caption(f"ทั้งหมด {len(filtered)} รายการ · หน้า {current_page}/{total_pages}")
    for offset in range(0, len(page_records), 3):
        columns = st.columns(3)
        for index, record in enumerate(page_records[offset : offset + 3]):
            with columns[index]:
                key_prefix = f"{record.get('__worksheet')}_{record.get('__source_row')}"
                _vendor_card(record, headers, admin_mode, key_prefix)
    previous_column, spacer, next_column = st.columns([1, 4, 1])
    with previous_column:
        if st.button("‹", disabled=current_page <= 1, use_container_width=True, key="vendor_prev"):
            st.session_state.vendor_page = current_page - 1
            st.rerun()
    with next_column:
        if st.button("›", disabled=current_page >= total_pages, use_container_width=True, key="vendor_next"):
            st.session_state.vendor_page = current_page + 1
            st.rerun()
