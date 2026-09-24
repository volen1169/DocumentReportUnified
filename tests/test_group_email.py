"""Focused Group E-mail schema and mailbox ownership rules."""

import datetime as dt
import ast
import io
from pathlib import Path
import unittest
from unittest.mock import Mock

import pandas as pd
from openpyxl import Workbook, load_workbook

from services.group_email import account_counts, group_email_payload, login_devices
from services.license_expiry import parse_expiration_date


HEADERS = ("Record ID", "Display Name", "Email", "Assigned Users", "Login Devices", "Company", "Expiry Date")


class GroupEmailTests(unittest.TestCase):
    def test_loaded_exact_fields_and_multiple_users_stay_on_one_row(self):
        frame = pd.DataFrame([{
            "Display Name": "PD", "Email": "PD@siamwin.co.th",
            "Assigned Users": "Parinda Putson / Suphakan Chotirak",
            "Login Devices": "2",
        }])
        self.assertEqual(account_counts(frame), (1, 1, 1))
        self.assertEqual(frame.iloc[0]["Assigned Users"], "Parinda Putson / Suphakan Chotirak")
        self.assertEqual(login_devices(frame.iloc[0]["Login Devices"]), 2)

    def test_blank_and_zero_device_counts(self):
        self.assertIsNone(login_devices(""))
        self.assertIsNone(login_devices(None))
        self.assertIsNone(login_devices("   "))
        self.assertEqual(login_devices("0"), 0)

    def test_positive_integer_device_count(self):
        self.assertEqual(login_devices("12"), 12)
        self.assertEqual(login_devices(2), 2)

    def test_invalid_device_counts(self):
        for value in ("-1", -1, "1.5", 1.5, "many", True, False):
            with self.subTest(value=value), self.assertRaises(ValueError):
                login_devices(value)

    def test_add_payload_uses_exact_keys_and_preserves_names(self):
        values = dict.fromkeys(HEADERS, "")
        values.update({"Display Name": "PD", "Email": "PD@siamwin.co.th",
                       "Assigned Users": "Parinda Putson / Suphakan Chotirak", "Login Devices": "2"})
        payload = group_email_payload(HEADERS, values)
        self.assertEqual(set(payload), set(HEADERS))
        self.assertEqual(payload["Login Devices"], 2)
        self.assertEqual(payload["Display Name"], "PD")
        self.assertEqual(payload["Assigned Users"], values["Assigned Users"])

    def test_edit_payload_allows_blank_devices_and_retains_unrelated_fields(self):
        values = {"Record ID": 17, "Display Name": "PD", "Email": "PD@siamwin.co.th",
                  "Assigned Users": "Parinda / Suphakan", "Login Devices": "", "Company": "SWI",
                  "Expiry Date": "2026-11-12 00:00:00"}
        payload = group_email_payload(HEADERS, values, ["other@siamwin.co.th"], original_email=values["Email"])
        self.assertIsNone(payload["Login Devices"])
        self.assertEqual(payload["Record ID"], 17)
        self.assertEqual(payload["Expiry Date"], values["Expiry Date"])

    def test_add_rejects_duplicate_email_regardless_of_case_or_outer_space(self):
        values = {"Email": " PD@siamwin.co.th ", "Login Devices": "0"}
        with self.assertRaisesRegex(ValueError, "1 Email"):
            group_email_payload(HEADERS, values, ["pd@SIAMWIN.co.th"])

    def test_missing_schema_and_blank_email_rejected(self):
        with self.assertRaisesRegex(ValueError, "Login Devices"):
            group_email_payload(("Email", "Assigned Users"), {"Email": "a@b.com"})
        with self.assertRaisesRegex(ValueError, "กรุณากรอก Email"):
            group_email_payload(HEADERS, {"Email": "  "})

    def test_blank_email_and_duplicates_excluded_from_account_total(self):
        frame = pd.DataFrame({"Email": ["  ", None, "PD@siamwin.co.th", " pd@SIAMWIN.co.th ", "other@example.com"]})
        self.assertEqual(account_counts(frame), (5, 3, 2))

    def test_expiry_parser_regression_for_loader_datetime_text(self):
        self.assertEqual(parse_expiration_date("2026-11-12 00:00:00"), dt.date(2026, 11, 12))
        self.assertEqual((dt.date(2026, 11, 12) - dt.date(2026, 9, 24)).days, 49)

    def test_add_optional_devices_and_exact_schema(self):
        headers = HEADERS + ("License Type", "Status")
        for raw, expected in (("", None), ("0", 0), ("2", 2)):
            with self.subTest(raw=raw):
                values = dict.fromkeys(headers, "")
                values.update({"Display Name": "PD", "Email": "PD@siamwin.co.th",
                               "Assigned Users": "Parinda / Suphakan", "Login Devices": raw,
                               "Company": "SWI", "License Type": "Business Basic",
                               "Expiry Date": "2026-11-12 00:00:00", "Status": "Active"})
                result = group_email_payload(headers, values)
                self.assertEqual(set(result), set(headers))
                self.assertEqual(result["Login Devices"], expected)
                self.assertEqual(result["Assigned Users"], "Parinda / Suphakan")
                self.assertEqual(account_counts(pd.DataFrame([result])), (1, 1, 1))

    def test_edit_device_transitions_and_other_fields(self):
        old = dict.fromkeys(HEADERS, "")
        old.update({"Record ID": 17, "Email": "PD@siamwin.co.th", "Assigned Users": "Parinda",
                    "Company": "SWI", "Expiry Date": "2026-11-12 00:00:00"})
        for before, after, expected in (("", "2", 2), ("2", "", None)):
            with self.subTest(before=before, after=after):
                existing = {**old, "Login Devices": before}
                edited = {**existing, "Assigned Users": "Parinda / Suphakan", "Login Devices": after}
                result = group_email_payload(HEADERS, edited, ["other@example.com"], old["Email"])
                self.assertEqual(result["Login Devices"], expected)
                self.assertEqual(result["Assigned Users"], "Parinda / Suphakan")
                self.assertEqual(result["Record ID"], existing["Record ID"])
                self.assertEqual(result["Company"], existing["Company"])
                self.assertEqual(result["Expiry Date"], existing["Expiry Date"])

    def test_exact_and_spaced_duplicate_email_rejected_on_add_and_edit(self):
        for new_email in ("PD@siamwin.co.th", " pd@SIAMWIN.co.th "):
            for original in (None, "original@example.com"):
                with self.subTest(new_email=new_email, original=original):
                    with self.assertRaisesRegex(ValueError, "1 Email"):
                        group_email_payload(HEADERS, {"Email": new_email},
                                            ["pd@siamwin.co.th"], original)

    def test_account_count_ignores_users_and_devices_without_rewriting_email(self):
        emails = [" ", "PD@siamwin.co.th", " pd@SIAMWIN.co.th "]
        frame = pd.DataFrame({"Email": emails, "Assigned Users": ["", "Parinda / Suphakan", "Other"],
                              "Login Devices": [None, 2, 5]})
        self.assertEqual(account_counts(frame), (3, 2, 1))
        self.assertEqual(frame["Email"].tolist(), emails)

    def test_real_loader_contract_excludes_non_account_sheets(self):
        # Execute the exact loader functions without importing Graph or accessing SharePoint.
        module = ast.parse(Path(__file__).resolve().parents[1].joinpath("services/excel_storage.py").read_text())
        functions = [node for node in module.body if isinstance(node, ast.FunctionDef)
                     and node.name in ("parse_password_sheet", "load_software_excels")]
        for function in functions:
            function.decorator_list = []
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append(list(HEADERS))
        sheet.append([1, "PD", "PD@siamwin.co.th", "Parinda / Suphakan", None, "SWI", "2026-11-12 00:00:00"])
        sheet.append([2, "", "  ", "", None, "", None])
        summary = workbook.create_sheet("Price Comparison")
        summary.append(["Annual License Price Comparison"])
        summary.append(["other data"])
        data = io.BytesIO()
        workbook.save(data)
        response = Mock(status_code=200, content=data.getvalue())
        get = Mock(side_effect=[Mock(json=lambda: {"id": "drive"}), response])
        scope = {"pd": pd, "io": io, "load_workbook": load_workbook,
                 "requests": Mock(get=get), "get_access_token": lambda: "placeholder",
                 "get_sp_site_id": lambda: "site", "GRAPH_URL": "https://example.invalid",
                 "SOFTWARE_FILE_MAP": {"Group Email": "Software_Group_Email.xlsx"},
                 "SHAREPOINT_FOLDER": "Update IT documents", "parse_password_sheet": None}
        exec(compile(ast.Module(body=functions, type_ignores=[]), "services/excel_storage.py", "exec"), scope)
        loaded, errors = scope["load_software_excels"]()
        self.assertEqual(errors, {})
        frame = loaded["Group Email"]
        self.assertEqual(account_counts(frame), (2, 1, 1))
        self.assertEqual(set(frame["Source Sheet"]), {"Sheet1"})
        self.assertEqual(frame.iloc[0]["Assigned Users"], "Parinda / Suphakan")
        self.assertIsNone(login_devices(frame.iloc[0]["Login Devices"]))

    def test_view_dialog_maps_all_fields_with_blank_device_placeholder(self):
        module = ast.parse(Path(__file__).resolve().parents[1].joinpath("DocumentReportUnified.py").read_text())
        function = next(node for node in module.body if isinstance(node, ast.FunctionDef)
                        and node.name == "view_group_email_record_dialog")
        function.decorator_list = []
        lines = []
        scope = {"st": Mock(text=lines.append), "parse_expiration_date": parse_expiration_date,
                 "_software_form_value": lambda value: "" if value is None else str(value)}
        exec(compile(ast.Module(body=[function], type_ignores=[]), "DocumentReportUnified.py", "exec"), scope)
        scope[function.name]({"Display Name": "PD", "Email": "PD@siamwin.co.th",
                              "Assigned Users": "Parinda / Suphakan", "Login Devices": None,
                              "Company": "SWI", "License Type": "Business Basic",
                              "Expiry Date": "2026-11-12 00:00:00", "Status": "Active"})
        self.assertEqual(lines, ["Display Name: PD", "Email: PD@siamwin.co.th",
                                 "Assigned Users: Parinda / Suphakan", "Login Devices: —",
                                 "Company: SWI", "License Type: Business Basic",
                                 "Expiry Date: 12 Nov 2026", "Status: Active"])


if __name__ == "__main__":
    unittest.main()
